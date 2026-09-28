"""Календарный план из Excel и CSV: шаблон, предпросмотр с ошибками по строкам, загрузка, замена прежнего плана."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from tests.conftest import login_as

pytestmark = pytest.mark.anyio


async def _new_site(client, admin) -> str:  # noqa: ANN001
    created = await client.post("/api/sites", headers=admin, json={"name": "Поликлиника на 300 посещений", "kind": "healthcare"})
    assert created.status_code == 201, created.text
    return created.json()["id"]


def _upload(name: str, content: bytes) -> dict:
    return {"file": (name, content, "application/octet-stream")}


async def test_template_is_filled_and_loaded_back(client):
    admin = await login_as(client, "admin")
    site = await _new_site(client, admin)
    response = await client.get(f"/api/sites/{site}/plan/template", headers=admin)
    assert response.status_code == 200 and response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "План" and {"Правила", "Пример", "Как заполнять"} <= set(book.sheetnames)
    plan = book["План"]
    assert [c.value for c in plan[1]] == [
        "Этап", "Работа", "Начало", "Окончание", "Правило «этап → техника»", "Вид работ по справочнику",
    ]  # fmt: skip
    assert any("E2" in str(v.sqref) for v in plan.data_validations.dataValidation)  # правило — из списка
    assert [r[0].value for r in book["Правила"].iter_rows(min_row=2)][:2] == ["Подготовка площадки", "Разработка котлована"]

    # заполняем, как заполнил бы человек: этап один раз на несколько работ, правило не везде
    for row in (
        ("Земляные работы", "Разработка котлована", date(2026, 10, 5), date(2026, 10, 30), None),
        (None, "Устройство дренажа", date(2026, 10, 26), date(2026, 11, 6), None),
        ("Нулевой цикл", "Фундаментная плита", date(2026, 11, 9), date(2026, 12, 4), "Бетонирование фундаментной плиты"),
    ):
        plan.append(list(row))
    filled = io.BytesIO()
    book.save(filled)

    preview = await client.post(f"/api/sites/{site}/plan/import", headers=admin, files=_upload("план.xlsx", filled.getvalue()))
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert (body["applied"], body["errors"], body["works"], body["existing"]) == (False, 0, 3, 0)
    earth, found = body["phases"]
    assert (earth["name"], earth["start"], earth["end"]) == ("Земляные работы", "2026-10-05", "2026-11-06")  # по работам
    pit, drain = earth["works"]
    assert pit["ruleKey"] == "excavation" and "по названию работы" in pit["notes"][0]
    assert drain["ruleKey"] is None and "не сверяем" in drain["notes"][0]
    assert found["works"][0]["ruleKey"] == "foundation_concrete"
    assert (await client.get(f"/api/stages?siteId={site}", headers=admin)).json() == []  # предпросмотр ничего не меняет

    applied = await client.post(
        f"/api/sites/{site}/plan/import?apply=true", headers=admin, files=_upload("план.xlsx", filled.getvalue())
    )
    assert applied.status_code == 200 and applied.json()["applied"] is True
    stages = (await client.get(f"/api/stages?siteId={site}", headers=admin)).json()
    assert [(s["level"], s["name"]) for s in stages] == [
        (1, "Земляные работы"), (2, "Разработка котлована"), (2, "Устройство дренажа"),
        (1, "Нулевой цикл"), (2, "Фундаментная плита"),
    ]  # fmt: skip
    assert stages[1]["parentId"] == stages[0]["id"] and stages[1]["ruleKey"] == "excavation"
    log = (await client.get("/api/audit?action=plan.import", headers=admin)).json()
    assert log and "этапов — 2, работ — 3" in log[0]["summary"]


async def test_windows_csv_with_errors_is_not_loaded(client):
    manager = await login_as(client, "manager")
    admin = await login_as(client, "admin")
    site = await _new_site(client, admin)
    csv = (
        "Этап;Работа;Начало;Окончание;Правило\n"
        "Подготовка;Подготовка площадки;01.10.2026;15.10.2026;\n"
        ";Ограждение;15.10.2026;01.10.2026;\n"  # окончание раньше начала
        "Каркас;Монтаж колонн;32.11.2026;20.12.2026;Монтаж каркаса\n"  # нет такой даты
        "Каркас;Монтаж плит;01.12.2026;20.12.2026;Сварка\n"  # нет такого правила
    ).encode("cp1251")  # так сохраняет CSV Excel в Windows
    preview = await client.post(f"/api/sites/{site}/plan/import", headers=manager, files=_upload("plan.csv", csv))
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["errors"] == 3
    prep, frame = body["phases"]
    assert prep["works"][1]["errors"] == ["Окончание раньше начала"] and prep["works"][1]["line"] == 3
    assert "не дата — «32.11.2026»" in frame["works"][0]["errors"][0]
    assert "Нет правила «Сварка»" in frame["works"][1]["errors"][0]
    refused = await client.post(f"/api/sites/{site}/plan/import?apply=true", headers=manager, files=_upload("plan.csv", csv))
    assert refused.status_code == 422 and "ошибки" in refused.json()["detail"]
    assert (await client.get(f"/api/stages?siteId={site}", headers=admin)).json() == []


async def test_replace_keeps_alert_history(client):
    admin = await login_as(client, "admin")
    csv = "Этап,Работа,Начало,Окончание\nЗемляные работы,Разработка котлована,2026-01-12,2027-03-01\n".encode()
    before = (await client.get("/api/alerts?siteId=s1", headers=admin)).json()
    assert any(a["stageId"] for a in before)
    preview = await client.post("/api/sites/s1/plan/import", headers=admin, files=_upload("plan.csv", csv))
    assert preview.json()["existing"] == 10  # в демо-плане ЖК 4 этапа и 6 работ — при замене удалятся
    done = await client.post("/api/sites/s1/plan/import?apply=true&replace=true", headers=admin, files=_upload("plan.csv", csv))
    assert done.status_code == 200, done.text
    stages = (await client.get("/api/stages?siteId=s1", headers=admin)).json()
    assert [s["name"] for s in stages] == ["Земляные работы", "Разработка котлована"]
    after = (await client.get("/api/alerts?siteId=s1", headers=admin)).json()
    assert len(after) == len(before) and all(a["stageId"] is None for a in after)  # история отклонений на месте


async def test_bad_files_and_access(client):
    admin, foreman = await login_as(client, "admin"), await login_as(client, "foreman")
    inspector = await login_as(client, "inspector")
    xls = b"\xd0\xcf\x11\xe0" + b"\x00" * 100
    old = await client.post("/api/sites/s1/plan/import", headers=admin, files=_upload("plan.xls", xls))
    assert old.status_code == 422 and ".xlsx" in old.json()["detail"]
    headerless = await client.post("/api/sites/s1/plan/import", headers=admin, files=_upload("plan.csv", b"a;b;c\n1;2;3\n"))
    assert headerless.status_code == 422 and "заголовок" in headerless.json()["detail"]
    for who in (foreman, inspector):  # план ведут руководитель проекта и администратор
        assert (await client.get("/api/sites/s1/plan/template", headers=who)).status_code == 403
        assert (await client.post("/api/sites/s1/plan/import", headers=who, files=_upload("p.csv", b"x"))).status_code == 403


async def test_catalog_column_by_number_or_name(client, monkeypatch):
    import httpx

    from app import mock_analytics
    from app.config import get_settings
    from app.services.analytics import runner

    monkeypatch.setattr(get_settings(), "deterministic_service_url", "http://analytics.test/deterministic")
    monkeypatch.setattr(runner, "_analytics", runner.Analytics(transport=httpx.ASGITransport(app=mock_analytics.app)))
    admin = await login_as(client, "admin")
    template = load_workbook(io.BytesIO((await client.get("/api/sites/s2/plan/template", headers=admin)).content))
    assert template["Виды работ"]["A2"].value.split(" — ")[0].isdigit()  # список «номер — название» для выбора

    csv = (
        "Этап;Работа;Начало;Окончание;Вид работ по справочнику\n"
        "Пристройка;Котлован пристройки;01.10.2026;20.10.2026;47\n"
        ";Выемка грунта;05.10.2026;25.10.2026;Выемка грунта котлована\n"  # по названию — оно в справочнике одно
        ";Стены подвала;26.10.2026;20.11.2026;Устройство ж/б конструкций\n"  # у школы таких видов работ два
        ";Подготовка;01.10.2026;03.10.2026;4 — Подготовка территории\n"  # сводный этап, не работа
    ).encode()
    preview = await client.post("/api/sites/s2/plan/import", headers=admin, files=_upload("p.csv", csv))
    works = preview.json()["phases"][0]["works"]
    assert [w["catalogStageId"] for w in works[:2]] == [47, 72]
    assert "укажите его номер" in works[2]["errors"][0]
    assert "сводный этап" in works[3]["errors"][0]
